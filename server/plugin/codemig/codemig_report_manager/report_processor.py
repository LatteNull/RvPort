# 标准库导入
import io
import os
import re
import shutil
import sqlite3
import sys
import tempfile
from urllib.parse import quote

# 第三方库导入
import sqlcipher3
from flask import Response
from openpyxl import Workbook
from sqlalchemy import case, func, union_all

# 项目路径设置
MODULE_PATH = os.path.realpath(
    os.path.join(os.path.dirname(__file__), '..', '..', '..'))
if MODULE_PATH not in sys.path:
    sys.path.append(MODULE_PATH)

# 项目导入 - common 模块
from common.db.SQLOperation import DatabaseException, db_session_scope, get_conn
from common.public import logger

# 项目导入 - framework 模块
from framework.log_manager.operation_constant import OpLogModule, OpLogResult
from framework.log_manager.operation_log import record_operation_log

# 项目导入 - plugin.codemig 模块
from plugin.codemig.codemig_common.common_git import generate_patch_from_initial_commit
from plugin.codemig.codemig_common.codemig_common_utils import (
    CodeTooLong,
    create_relative_file_path,
    get_database_path,
    USE_ABSOLUTE_PATH,
)
from plugin.codemig.codemig_common.migration_class import MigrationClass
from plugin.codemig.codemig_task_manager.query_processor import (
    query_dir_id_by_task_id,
    query_task_name_from_database,
)
from plugin.codemig.dependence_process.models import (
    CodemigTaskDependence,
    DependencePackages,
    FileDependence,
    FileDependenceSuggestion,
)
from plugin.codemig.solver.solvers_mapping import SolverType

LOGGER = logger.init_log('codemig_report_manager_request_processor',
                         'codemig.log')

TASK_NAME_SQL = ("SELECT task_name FROM codemig_tasks  WHERE task_id = ?")

# 源码扫描按问题类型做统计（在用户特定数据库中查询）
TASK_INFO_SQL = (
    "SELECT"
    "    Suggestion.type AS suggestion_type,"
    "    COUNT(Suggestion.change_id) AS num_changes,"
    "    SUM(Suggestion.end_line_number - Suggestion.start_line_number + 1) AS num_changed_lines "
    "FROM codemig_task_directory_file AS File "
    "JOIN codemig_task_directory_file_suggestion AS Suggestion "
    "    ON File.file_id = Suggestion.file_id "
    "WHERE File.has_removed = 0 "
    "    AND File.type = 0 "
    "GROUP BY Suggestion.type "
    "ORDER BY Suggestion.type ASC")

# 除源码扫描外，其他类型的统计（在用户特定数据库中查询）
TASK_FILE_INFO_SQL = (
    "SELECT"
    "    File.type AS file_type,"
    "    COUNT(File.file_id) AS number_files "
    "FROM codemig_task_directory_file AS File "
    "WHERE File.type > 0 "
    "GROUP BY File.type "
    "ORDER BY File.type ASC")

TASK_CHANGE_SQL = (
    "SELECT"
    "    File.file_path AS file_path,"
    "    Suggestion.type AS suggestion_type,"
    "    Suggestion.start_line_number AS start_line_number,"
    "    Suggestion.end_line_number AS end_line_number,"
    "    Suggestion.suggestion AS suggestion,"
    "    Suggestion.code_too_long AS code_too_long,"
    "    Suggestion.code AS code "
    "FROM codemig_task_directory_file AS File "
    "JOIN codemig_task_directory_file_suggestion AS Suggestion "
    "    ON File.file_id = Suggestion.file_id "
    "WHERE File.has_removed = 0 "
    "    AND File.type = 0 "
    "ORDER BY File.file_path, Suggestion.type, Suggestion.start_line_number, "
    "Suggestion.end_line_number ASC"
)

TASK_LIBRARY_SQL = (
    "SELECT"
    "    File.file_path AS file_path,"
    "    GROUP_CONCAT(SoftwareSuggestion.support_architecture) AS support_architecture,"
    "    GROUP_CONCAT(SoftwareSuggestion.suggestion) AS suggestion "
    "FROM codemig_task_directory_file AS File "
    "JOIN codemig_task_directory_software_suggestion AS SoftwareSuggestion "
    "    ON File.file_id = SoftwareSuggestion.file_id "
    "WHERE File.has_removed = 0 "
    "    AND File.type = 1 "
    "GROUP BY File.file_id "
    "ORDER BY File.file_path ASC")

MISSING_FILE_SQL = (
    "SELECT File.file_path AS missing_file, GROUP_CONCAT(ExistFile.architecture_path) "
    "AS exist_file_path "
    "FROM codemig_task_directory_file AS File "
    "JOIN codemig_task_directory_existing_architecture AS ExistFile "
    "    ON File.file_id = ExistFile.file_id "
    "WHERE File.has_removed = 0 AND File.type = 2 "
    "GROUP BY File.file_id "
    "ORDER BY File.file_path ASC")




def convert_task_info(task_info, task_name):
    rows = []
    for solver_type, num_changes, num_lines in task_info:
        suggestion_desc = SolverType(solver_type).description
        rows.append((task_name, suggestion_desc, num_changes, num_lines))
    return rows


def convert_task_file_info(task_file_info, task_name):
    rows = []
    for file_type, number_files in task_file_info:
        file_type_desc = MigrationClass(file_type).description
        rows.append((task_name, file_type_desc, number_files))
    return rows


def convert_change_info(change_info):
    rows = []
    for (file_path, solver_type, start_lineno, end_lineno, suggestion,
         code_too_long, code) in change_info:
        suggestion_desc = SolverType(solver_type).description
        file_name = os.path.basename(file_path)
        if not USE_ABSOLUTE_PATH:
            file_path = create_relative_file_path(file_path)
        if code_too_long is not None and CodeTooLong(code_too_long) != CodeTooLong.HAS_SOURCE_CODE:
            code = '源代码过长，无法显示。'
        rows.append(
            (file_name, file_path, suggestion_desc, start_lineno, end_lineno, suggestion, code))
    return rows


def convert_library_info(library_info):
    rows = []
    for file_path, support_architecture, suggestion in library_info:
        file_name = os.path.basename(file_path)
        if not USE_ABSOLUTE_PATH:
            file_path = create_relative_file_path(file_path)
        rows.append((file_name, file_path, support_architecture, suggestion))
    return rows


def convert_missing_file_info(missing_file_info):
    rows = []
    for (file_path, exist_path) in missing_file_info:
        file_name = os.path.basename(file_path)
        if not USE_ABSOLUTE_PATH:
            file_path = create_relative_file_path(file_path)
            others = exist_path.split(',')
            others = [create_relative_file_path(other) for other in others]
            exist_path = ','.join(others)
        rows.append((file_name, file_path, exist_path))
    return rows


def save_excel_data(report_info, task_name):
    (task_info, task_file_info, change_info, library_info, missing_file_info,
     dependence_info) = report_info
    # 创建一个新的工作簿
    wb = Workbook()

    # 重命名第一个工作表
    ws1 = wb.active
    ws1.title = "任务信息"
    ws1.append(("任务名称", "建议类型", "问题数量", "影响行数"))

    # 逐行写入数据
    for row in convert_task_info(task_info, task_name):
        ws1.append(row)
    for row in convert_task_file_info(task_file_info, task_name):
        ws1.append(row)

    # 创建第二个工作表
    ws2 = wb.create_sheet(title="源代码信息")
    ws2.append(("文件名称", "文件路径", "建议类型", "起始行", "终止行", "建议", "源代码"))

    # 逐行写入数据
    for row in convert_change_info(change_info):
        ws2.append(row)

    # 创建第三个工作表
    ws3 = wb.create_sheet(title="软件包信息")
    ws3.append(("文件名称", "文件路径", "支持架构", "建议信息"))

    # 逐行写入数据
    for row in convert_library_info(library_info):
        ws3.append(row)

    # 创建第四个工作表
    ws4 = wb.create_sheet(title="缺失文件信息")
    ws4.append(("文件名称", "文件路径", "现有架构文件"))

    # 逐行写入数据
    for row in convert_missing_file_info(missing_file_info):
        ws4.append(row)

    if dependence_info and len(dependence_info) > 0:
        # 创建第五个工作表
        ws4 = wb.create_sheet(title="依赖库信息")
        ws4.append(("依赖包", "版本号", "问题个数"))
        # 逐行写入数据
        for row in dependence_info:
            ws4.append(row)
    # 将工作簿保存到内存中的字节流
    excel_data = io.BytesIO()
    wb.save(excel_data)
    excel_data.seek(0)
    return excel_data


def query_report_info_from_database(user_name, user_id, task_id):
    LOGGER.info(f"Querying the report data for user {user_id} task_id {task_id}.")

    status, task_name = query_task_name_from_database(task_id)
    if not status or task_name is None:
        LOGGER.error(f"Failed to get task name for task_id {task_id}")
        return None, None

    # 获取 dir_id
    status, dir_id = query_dir_id_by_task_id(task_id)
    if not status or dir_id is None:
        LOGGER.error(f"Failed to get dir_id for task_id {task_id}")
        return None, None

    # 从用户特定的数据库查询文件相关信息
    user_db_path = get_database_path(user_name, dir_id)
    if not user_db_path.exists():
        LOGGER.error(f"Database {user_db_path} not found")
        return None, None

    user_conn = None
    try:
        user_conn = get_conn(user_db_path)
        cursor = user_conn.cursor()

        # 查询源码扫描统计信息
        cursor.execute(TASK_INFO_SQL)
        source_code_summary_info = cursor.fetchall()

        # 查询其他类型文件统计信息
        cursor.execute(TASK_FILE_INFO_SQL)
        file_summary_info = cursor.fetchall()

        # 查询源代码详细信息
        cursor.execute(TASK_CHANGE_SQL)
        change_info = cursor.fetchall()

        # 查询软件包信息
        cursor.execute(TASK_LIBRARY_SQL)
        library_info = cursor.fetchall()

        # 查询缺失文件信息
        cursor.execute(MISSING_FILE_SQL)
        missing_file_info = cursor.fetchall()

    except (sqlite3.Error, sqlcipher3.dbapi2.DatabaseError) as err:
        LOGGER.error("Database operation failed: %s", err)
        raise
    finally:
        if user_conn:
            user_conn.close()

    # 查询依赖信息（使用 ORM，从主数据库）
    dependence_info = query_dependence_info(task_id)

    return (task_name, (source_code_summary_info, file_summary_info, change_info, library_info,
                            missing_file_info, dependence_info))


def query_dependence_info(task_id):
    """
    查询依赖信息
    Args:
        task_id: 任务ID
    Returns:
        list: 依赖信息列表，每个元素为 (dependence_package, version, total_count)
    """
    try:
        with db_session_scope() as session:
            # 第一部分：统计 file_dependence_suggestion 的数量
            query1 = session.query(
                DependencePackages.dependence_package,
                DependencePackages.dependence_id,
                DependencePackages.version,
                func.count(FileDependenceSuggestion.id).label('count_value')
            ).join(
                CodemigTaskDependence,
                DependencePackages.dependence_id == CodemigTaskDependence.dependence_id
            ).outerjoin(
                FileDependenceSuggestion,
                FileDependenceSuggestion.dependence_id == CodemigTaskDependence.dependence_id
            ).filter(
                CodemigTaskDependence.task_id == task_id
            ).group_by(
                CodemigTaskDependence.dependence_id,
                DependencePackages.dependence_package,
                DependencePackages.version
            )
            
            # 第二部分：统计 file_dependence 中 type > 0 的文件数量
            query2 = session.query(
                DependencePackages.dependence_package,
                DependencePackages.dependence_id,
                DependencePackages.version,
                func.sum(
                    case(
                        ((FileDependence.file_id.isnot(None)) & (FileDependence.type > 0), 1),
                        else_=0
                    )
                ).label('count_value')
            ).join(
                CodemigTaskDependence,
                DependencePackages.dependence_id == CodemigTaskDependence.dependence_id
            ).outerjoin(
                FileDependence,
                FileDependence.dependence_id == CodemigTaskDependence.dependence_id
            ).filter(
                CodemigTaskDependence.task_id == task_id
            ).group_by(
                CodemigTaskDependence.dependence_id,
                DependencePackages.dependence_package,
                DependencePackages.version
            )
            
            # 使用 UNION ALL 合并两个查询结果
            # 需要确保两个查询的列顺序一致：dependence_package, dependence_id, version, count_value
            union_stmt = union_all(query1, query2).subquery()
            
            # 外层查询：按 dependence_package 和 version 分组，求和
            # 使用列索引访问，因为 union_all 后列名可能不一致
            final_query = session.query(
                union_stmt.c[0].label('dependence_package'),  # dependence_package
                union_stmt.c[2].label('version'),  # version
                func.sum(union_stmt.c[3]).label('total_count')  # count_value
            ).group_by(
                union_stmt.c[0],  # dependence_package
                union_stmt.c[2]   # version
            ).order_by(
                func.sum(union_stmt.c[3]).desc()  # count_value
            )
            
            results = final_query.all()
            return [
                (result.dependence_package, result.version, result.total_count or 0)
                for result in results
            ]
    except (DatabaseException, OSError, IOError) as err:
        LOGGER.error("Database operation failed: %s", err)
        raise


def create_report_response(user_name, task_id, request):
    user_id = request.user_id
    try:
        (task_name, report_info) = query_report_info_from_database(user_name, user_id, task_id)
        if task_name is None:
            record_operation_log(request.user_id, request.remote_addr, OpLogModule.CODEMIG,
                                 'Download Report', OpLogResult.FAILED,
                                 f'No task found of {user_name}: {task_id}.')
            return {"code": -1, "msg": f"No task found"}
        quote_safe_name = quote(re.sub(r'[/\\:*?"<>|\x00-\x1F ]', '_', task_name))
        excel_data = save_excel_data(report_info, task_name)
    except  (sqlite3.Error, sqlcipher3.dbapi2.DatabaseError) as err:
        record_operation_log(request.user_id, request.remote_addr, OpLogModule.CODEMIG,
                             'Download Report', OpLogResult.FAILED,
                             f'Download report of  {task_id} failed: database operation failed.')
        return {"code": -1, "msg": f"Database operation failed:{str(err)}"}
    except (IndexError, AttributeError, ValueError, TypeError, PermissionError, IOError,
            RuntimeError) as e:
        failed_detail = f"Preparing data failed:{str(e)}"
        record_operation_log(request.user_id, request.remote_addr, OpLogModule.CODEMIG,
                             'Download Report', OpLogResult.FAILED,
                             f'Download report of  {task_id} failed. {failed_detail}.')
        LOGGER.error(failed_detail, exc_info=True)
        return {"code": -1, "msg": failed_detail}
    response = Response(excel_data,
                        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response.headers['Content-Disposition'] = f"attachment; filename*=UTF-8''{quote_safe_name}.xlsx"
    record_operation_log(request.user_id, request.remote_addr, OpLogModule.CODEMIG,
                         'Download Report', OpLogResult.SUCCESS, f'Download report of {task_name}.')
    return response


def generate_patch_file(user_name, task_id):
    """
    使用 GitPython 生成从初始提交到当前版本的 patch 文件
    Args:
        user_name: 用户名
        task_id: 任务ID
    Returns:
        tuple: (patch_file_path, error_message)
            patch_file_path: 如果成功，返回临时 patch 文件路径；如果失败，返回 None
            error_message: 如果失败，返回错误信息；如果成功，返回 None
    """
    LOGGER.info(f"Generating patch file for user {user_name} task_id {task_id}.")
    try:
        # 获取任务信息
        status, task_name = query_task_name_from_database(task_id)
        if not status or task_name is None:
            LOGGER.error(f"Failed to get task name for task_id {task_id}")
            return None, "Task not found"

        # 获取任务的目录信息
        status, dir_id = query_dir_id_by_task_id(task_id)
        if not status or dir_id is None:
            LOGGER.error(f"Failed to get dir_id for task_id {task_id}")
            return None, "Directory not found"

        # 使用 GitPython 生成 patch
        status, patch_content = generate_patch_from_initial_commit(user_name, dir_id)
        if not status:
            LOGGER.error(f"Failed to generate patch: {patch_content}")
            return None, patch_content

        # 创建临时文件保存 patch
        temp_dir = tempfile.mkdtemp()
        patch_file = os.path.join(temp_dir, f"{task_name}.patch")
        with open(patch_file, 'w', encoding='utf-8') as f:
            f.write(patch_content)
        LOGGER.info(f"Successfully generated patch file: {patch_file}")
        return patch_file, None

    except (sqlite3.Error, sqlcipher3.dbapi2.DatabaseError, RuntimeError) as e:
        LOGGER.error(f"Database error: {e}")
        return None, "Database error"
    except (OSError, PermissionError, FileNotFoundError, FileExistsError) as e:
        LOGGER.error(f"File system error: {e}")
        return None, "File system error"
    except UnicodeDecodeError as e:
        LOGGER.error(f"Encoding error while reading files: {e}")
        return None, "File encoding error"



def create_patch_response(user_name, task_id, request):
    """
    创建patch文件下载响应
    """
    try:
        patch_file, error = generate_patch_file(user_name, task_id)
        if patch_file is None:
            record_operation_log(request.user_id, request.remote_addr, OpLogModule.CODEMIG,
                                 'Download Patch', OpLogResult.FAILED, f'Generate patch failed for {user_name}: {task_id}. Error: {error}')
            return {"code": 1, "msg": f"Generate patch failed: {error}"}

        # 读取patch文件内容
        with open(patch_file, 'rb') as f:
            patch_data = f.read()

        # 清理临时文件
        shutil.rmtree(os.path.dirname(patch_file))

        # 获取任务名称用于文件名
        status, task_name = query_task_name_from_database(task_id)
        if not status or task_name is None:
            LOGGER.error(f"Failed to get task name for task_id {task_id}")
            return {"code": 1, "msg": "Task not found"}
        quote_safe_name = quote(re.sub(r'[/\\:*?"<>|\x00-\x1F ]', '_', task_name))

        response = Response(patch_data, mimetype="text/plain")
        response.headers['Content-Disposition'] = f"attachment; filename*=UTF-8''{quote_safe_name}.patch"
        record_operation_log(request.user_id, request.remote_addr, OpLogModule.CODEMIG, 'Download Patch', OpLogResult.SUCCESS, f'Download patch of {task_name}.')
        return response
    except (FileNotFoundError, PermissionError, OSError, UnicodeDecodeError, sqlite3.Error, sqlcipher3.dbapi2.DatabaseError) as e:
        failed_detail = f"Generate patch failed: {str(e)}"
        record_operation_log(request.user_id, request.remote_addr, OpLogModule.CODEMIG,
                             'Download Patch', OpLogResult.FAILED, f'Download patch of {task_id} failed. {failed_detail}.')
        LOGGER.error(failed_detail, exc_info=True)
        return {"code": 1, "msg": failed_detail}
